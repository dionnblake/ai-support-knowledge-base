import docx
from pathlib import Path
from typing import List
from pypdf import PdfReader
from langchain_core.documents import Document

def load_pdf(file_path: Path) -> List[Document]:
    """
    Loads text from a PDF file page by page.
    
    Args:
        file_path (Path): Path to the PDF file.
        
    Returns:
        List[Document]: List of LangChain Document objects (one per page).
    """
    documents = []
    try:
        reader = PdfReader(file_path)
        for page_idx, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                # Clean up any bad byte sequences
                cleaned_text = text.encode("utf-8", errors="ignore").decode("utf-8")
                documents.append(Document(
                    page_content=cleaned_text,
                    metadata={
                        "source": str(file_path),
                        "page": page_idx  # 0-indexed page number
                    }
                ))
    except Exception as e:
        raise ValueError(f"Error reading PDF file {file_path.name}: {str(e)}")
        
    return documents

def load_docx(file_path: Path) -> List[Document]:
    """
    Loads text from a Word (.docx) document.
    
    Args:
        file_path (Path): Path to the DOCX file.
        
    Returns:
        List[Document]: List containing a single Document object with the full text.
    """
    try:
        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        full_text = "\n".join(paragraphs)
        
        # Clean up encoding issues
        cleaned_text = full_text.encode("utf-8", errors="ignore").decode("utf-8")
        
        return [Document(
            page_content=cleaned_text,
            metadata={
                "source": str(file_path),
                "page": 0  # DOCX doesn't have native pages in text extraction
            }
        )]
    except Exception as e:
        raise ValueError(f"Error reading DOCX file {file_path.name}: {str(e)}")

def load_txt(file_path: Path) -> List[Document]:
    """
    Loads text from a plain text (.txt) file.
    
    Args:
        file_path (Path): Path to the TXT file.
        
    Returns:
        List[Document]: List containing a single Document object.
    """
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
            
        return [Document(
            page_content=text,
            metadata={
                "source": str(file_path),
                "page": 0
            }
        )]
    except Exception as e:
        raise ValueError(f"Error reading TXT file {file_path.name}: {str(e)}")

def load_document(file_path: Path) -> List[Document]:
    """
    Dispatches file loading to the appropriate loader based on extension.
    
    Args:
        file_path (Path): Path to the file.
        
    Returns:
        List[Document]: Parsed LangChain Document objects.
    """
    suffix = file_path.suffix.lower()
    
    if suffix == ".pdf":
        return load_pdf(file_path)
    elif suffix == ".docx":
        return load_docx(file_path)
    elif suffix in (".txt", ".log", ".md"):
        return load_txt(file_path)
    else:
        raise ValueError(f"Unsupported file format: {suffix}")
